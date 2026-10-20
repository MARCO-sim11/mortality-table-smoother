"""Module 6 - exporter.py. Run:   python -m pytest tests/test_exporter.py -v

Each format is read back with an independent reader and compared with what went in:
  * CSV with pandas.read_csv (both separators, both decimal marks), exactly;
  * the HMD-style text with numpy.loadtxt, to the precision printed, and line by line for the layout;
  * the workbook with openpyxl: sheet names and order, the cells, the embedded image;
  * the chart by decoding the PNG and looking for the colours of the zones.
"""
from __future__ import annotations

import importlib
import io
from datetime import date

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from PIL import Image

import synthetic as S
from src import config as C
from src import evaluator as EV
from src import exporter as X
from src import joiner as J
from src.spline_all import fit_spline_all

SEL = {"zone1": "poly3", "zone2": "makeham", "zone3": "kannisto"}


@pytest.fixture(scope="module")
def case():
    t = S.make_sex_table("Female")                                         # ages 0..100, Makeham + infant term
    df = pd.DataFrame({"sex": "Female", "age": t["age"], "Dx": t["Dx"].astype(float), "Ex": t["Ex"], "mx_raw": t["Dx"] / t["Ex"]})
    cfg = C.default_config(df["age"].unique())
    fits = {}
    for m in C.METHODS.values():
        sub = C.zone_slice(df, cfg, m.zone)
        fit = getattr(importlib.import_module(f"src.{m.module}"), m.fit_name)
        fits[m.key] = fit(sub["age"].to_numpy(), sub["Dx"].to_numpy(), sub["Ex"].to_numpy(), cfg)
    mu = J.join_zones(*(fits[SEL[z]] for z in C.ZONES), cfg)
    return {"df": df, "cfg": cfg, "fits": fits, "mu": mu, "table": J.build_life_table(mu),
            "criteria": EV.evaluate_fits(fits, df, cfg).rename(columns=EV.CRITERIA_LABELS)}


# --------------------------------------------------------------------------
# Step 1: columns
# --------------------------------------------------------------------------


def test_select_columns_keeps_the_table_order_and_always_the_age(case):
    t = case["table"]
    assert list(X.select_columns(t).columns) == X.TABLE_COLUMNS == ["age", "mu", "qx", "px", "lx", "dx", "ex"]
    assert list(X.select_columns(t, ["ex", "qx"]).columns) == ["age", "qx", "ex"]            # age added, table order
    assert list(X.select_columns(t, ["age", "dx", "dx"]).columns) == ["age", "dx"]
    sub = X.select_columns(t, ["mu"])
    pd.testing.assert_series_equal(sub["mu"], t["mu"])
    assert sub is not t and X.select_columns(t) is not t                                      # a copy


def test_select_columns_refuses_unknown_names_and_an_empty_choice(case):
    with pytest.raises(ValueError, match="inconnue.*Lx.*colonnes disponibles : age, mu"):
        X.select_columns(case["table"], ["mu", "Lx"])
    with pytest.raises(ValueError, match="au moins une colonne"):
        X.select_columns(case["table"], [])
    with pytest.raises(ValueError, match="au moins une colonne"):
        X.select_columns(case["table"], ["age"])


# --------------------------------------------------------------------------
# Steps 2 and 3: CSV and HMD text
# --------------------------------------------------------------------------


@pytest.mark.parametrize("sep,decimal", [(",", "."), (";", "."), (";", ",")])
def test_csv_reads_back_exactly_with_every_separator_and_decimal_mark(case, sep, decimal):
    text = X.to_csv(case["table"], sep=sep, decimal=decimal)
    back = pd.read_csv(io.StringIO(text), sep=sep, decimal=decimal, float_precision="round_trip")
    pd.testing.assert_frame_equal(back, case["table"].astype({"age": "int64"}), check_dtype=True)
    assert text.splitlines()[0] == sep.join(X.TABLE_COLUMNS) and "\r" not in text and text.endswith("\n")
    assert len(text.splitlines()) == 1 + len(case["table"])


def test_csv_column_choice_and_the_decimal_comma_in_the_text(case):
    text = X.to_csv(case["table"], ["qx", "lx"], sep=";", decimal=",")
    assert text.splitlines()[0] == "age;qx;lx"
    row = text.splitlines()[1].split(";")
    assert row[0] == "0" and "," in row[1] and "." not in row[1]
    assert float(row[1].replace(",", ".")) == case["table"]["qx"].iloc[0]


@pytest.mark.parametrize("sep,decimal,match", [("\t", ".", "Séparateur"), (" ", ".", "Séparateur"), (";", ";", "Marque décimale"),
                                               (";", "'", "Marque décimale"), (",", ",", "identiques")])
def test_csv_refuses_a_bad_separator_or_decimal_or_the_same_comma_twice(case, sep, decimal, match):
    with pytest.raises(ValueError, match=match):
        X.to_csv(case["table"], sep=sep, decimal=decimal)


def test_hmd_text_has_the_hmd_layout_and_reads_back_to_the_printed_precision(case):
    t = case["table"]
    text = X.to_hmd_txt(t, title="France", year=2020, modified=date(2026, 10, 19))
    lines = text.split("\n")
    assert lines[0] == "France, smoothed life table (period 1x1)\tLast modified: 19 Oct 2026, Mortality Table Smoother"
    assert lines[1] == "" and text.endswith("\n") and len(lines) == 3 + len(t) + 1
    assert lines[2].split() == ["Year", "Age", "mu", "qx", "px", "lx", "dx", "ex"]
    assert len({len(line) for line in lines[2:-1]}) == 1                                     # fixed-width: every line as long
    data = np.loadtxt(io.StringIO(text), skiprows=3)
    assert data.shape == (len(t), 8) and np.all(data[:, 0] == 2020) and np.array_equal(data[:, 1], t["age"])
    for j, col in enumerate(["mu", "qx", "px"], start=2):
        np.testing.assert_allclose(data[:, j], t[col], atol=6e-13)
    for j, col in enumerate(["lx", "dx", "ex"], start=5):
        np.testing.assert_allclose(data[:, j], t[col], atol=6e-5)


def test_hmd_text_for_a_generation_table_a_chosen_subset_and_an_unknown_year(case):
    t = case["table"]
    cohort = X.to_hmd_txt(t, ["qx", "ex"], title="Luxembourg", cohort=1950, modified=date(2026, 1, 2))
    assert cohort.split("\n")[0].startswith("Luxembourg, smoothed life table (cohort 1x1)\tLast modified: 02 Jan 2026")
    assert cohort.split("\n")[2].split() == ["Cohort", "Age", "qx", "ex"]
    assert np.loadtxt(io.StringIO(cohort), skiprows=3).shape == (len(t), 4)
    no_year = X.to_hmd_txt(t, ["mu"], modified=date(2026, 1, 2))
    assert no_year.split("\n")[0].startswith("smoothed life table (period 1x1)") and no_year.split("\n")[2].split() == ["Age", "mu"]
    assert np.loadtxt(io.StringIO(no_year), skiprows=3).shape == (len(t), 2)


def test_hmd_text_keeps_seven_digits_for_a_tiny_rate(case):
    t = pd.DataFrame({"age": [8], "mu": [3.14159e-5], "qx": [3.14159e-5], "px": [0.99996858], "lx": [100000.0], "dx": [3.1416], "ex": [80.0]})
    assert "0.000031415900" in X.to_hmd_txt(t, modified=date(2026, 1, 2))


# --------------------------------------------------------------------------
# The tables of the Excel sheets
# --------------------------------------------------------------------------


def test_raw_table_is_age_dx_ex_mx_raw_in_age_order(case):
    shuffled = case["df"].sample(frac=1.0, random_state=1)
    raw = X.raw_table(shuffled)
    assert list(raw.columns) == ["age", "Dx", "Ex", "mx_raw"] and list(raw["age"]) == list(range(101))
    np.testing.assert_array_equal(raw["Dx"], case["df"]["Dx"])
    assert list(raw.index) == list(range(101))


def test_parameter_table_lists_every_parameter_by_zone_and_method_and_marks_the_chosen_ones(case):
    p = X.parameter_table(case["fits"], SEL)
    assert list(p.columns) == ["Zone", "Méthode", "Clé", "Paramètre", "Valeur", "Retenue"]
    assert list(dict.fromkeys(p["Clé"])) == list(C.METHODS)                                   # registry order
    poly = p[p["Clé"] == "poly3"].set_index("Paramètre")["Valeur"]
    assert list(poly.index) == ["a0", "a1", "a2", "a3", "n_params"]
    for k in range(4):
        assert poly[f"a{k}"] == case["fits"]["poly3"].params[f"a{k}"]
    knots = p[(p["Clé"] == "spline") & p["Paramètre"].str.startswith("knots[")]
    assert list(knots["Valeur"]) == case["fits"]["spline"].params["knots"]
    assert set(p.loc[p["Retenue"], "Clé"]) == set(SEL.values()) and p["Zone"].isin(["Zone I", "Zone II", "Zone III"]).all()
    assert not X.parameter_table(case["fits"])["Retenue"].any()
    only = X.parameter_table({"kannisto": case["fits"]["kannisto"]}, SEL)
    assert set(only["Clé"]) == {"kannisto"} and set(only["Paramètre"]) == {"a", "b", "x0", "n_params"} and only["Retenue"].all()


def test_method_curves_give_each_tested_method_its_own_curve_and_q(case):
    curves = X.method_curves(case["fits"])
    assert list(curves) == list(C.METHODS) and all(list(c.columns) == ["age", "mu", "qx"] for c in curves.values())
    k = curves["kannisto"]
    assert k["age"].iloc[0] == 85 and k["age"].iloc[-1] == 110                                # zone III runs to omega3
    np.testing.assert_allclose(k["mu"], case["fits"]["kannisto"].mu_hat, rtol=0)
    np.testing.assert_allclose(k["qx"], 1.0 - np.exp(-k["mu"]), rtol=1e-12)
    assert list(X.method_curves({"gompertz": case["fits"]["gompertz"]})) == ["gompertz"]


# --------------------------------------------------------------------------
# The chart
# --------------------------------------------------------------------------


def close_to(image: np.ndarray, hex_color: str, tol: int = 12, alpha: float = 1.0) -> int:
    """Pixels within `tol` of the colour as drawn over white with opacity `alpha` (the crude points are drawn at 0.75)."""
    rgb = np.array([int(hex_color[i:i + 2], 16) for i in (1, 3, 5)])
    drawn = alpha * rgb + (1.0 - alpha) * 255.0
    return int(np.sum(np.all(np.abs(image[:, :, :3].astype(float) - drawn) <= tol, axis=2)))


def test_the_chart_is_a_png_with_a_curve_in_each_zone_colour_and_the_crude_points(case):
    png = X.chart_image(X.raw_table(case["df"]), case["mu"], case["cfg"], "Female", "Titre de test", dpi=100)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    img = Image.open(io.BytesIO(png))
    assert img.size == (900, 540)                                                             # 9 x 5.4 inches at 100 dpi
    arr = np.array(img.convert("RGB"))
    for zone, colour in X.ZONE_COLORS.items():
        assert close_to(arr, colour) > 150, zone                                              # a visible line in each zone colour
    assert close_to(arr, X.SEX_COLORS["Female"], alpha=0.75) > 100                          # the crude points
    male = np.array(Image.open(io.BytesIO(X.chart_image(X.raw_table(case["df"]), case["mu"], case["cfg"], "Male", dpi=100))).convert("RGB"))
    assert close_to(male, X.SEX_COLORS["Male"], alpha=0.75) > 100 and close_to(male, X.SEX_COLORS["Female"], alpha=0.75) < 30


def test_the_chart_copes_with_a_curve_that_starts_late_and_with_zero_crude_rates(case):
    mu = case["mu"].copy()
    mu[:20] = np.nan                                                          # a table that starts at 20
    raw = X.raw_table(case["df"])
    raw.loc[raw["age"] == 50, "mx_raw"] = 0.0
    png = X.chart_image(raw, mu, case["cfg"], "Female", dpi=60)
    assert Image.open(io.BytesIO(png)).size == (540, 324)


# --------------------------------------------------------------------------
# Steps 4 and 5: the workbook
# --------------------------------------------------------------------------


def workbook(case, **kw):
    return load_workbook(io.BytesIO(X.build_excel(case["table"], X.raw_table(case["df"]), X.parameter_table(case["fits"], SEL),
                                                  case["criteria"], **kw)))


def sheet_frame(ws) -> pd.DataFrame:
    rows = list(ws.values)
    return pd.DataFrame(rows[1:], columns=rows[0])


def test_the_workbook_has_the_four_sheets_of_the_spec_in_order(case):
    wb = workbook(case)
    assert wb.sheetnames == ["Table_lissée", "Taux_bruts", "Paramètres", "Critères"]
    table = sheet_frame(wb["Table_lissée"])
    pd.testing.assert_frame_equal(table, case["table"].astype({"age": "int64"}), check_dtype=False, rtol=1e-14)
    raw = sheet_frame(wb["Taux_bruts"])
    assert list(raw.columns) == ["age", "Dx", "Ex", "mx_raw"] and len(raw) == 101
    np.testing.assert_allclose(raw["Dx"], case["df"]["Dx"])
    params = sheet_frame(wb["Paramètres"])
    assert list(params.columns) == ["Zone", "Méthode", "Clé", "Paramètre", "Valeur", "Retenue"] and len(params) == len(X.parameter_table(case["fits"], SEL))
    crit = sheet_frame(wb["Critères"])
    assert list(crit.columns) == list(case["criteria"].columns) and len(crit) == 9
    np.testing.assert_allclose(crit["AIC"], case["criteria"]["AIC"], rtol=1e-12)


def test_the_workbook_header_is_bold_the_first_row_frozen_and_the_columns_wide_enough(case):
    ws = workbook(case)["Table_lissée"]
    assert all(c.font.bold for c in ws[1]) and not ws["A2"].font.bold
    assert ws.freeze_panes == "A2"
    assert all(ws.column_dimensions[c].width >= 8 for c in "ABCDEFG")


def test_the_column_choice_applies_to_the_table_sheet_only(case):
    wb = workbook(case, columns=["qx", "ex"])
    assert [c.value for c in wb["Table_lissée"][1]] == ["age", "qx", "ex"]
    assert [c.value for c in wb["Taux_bruts"][1]] == ["age", "Dx", "Ex", "mx_raw"]


def test_the_chart_sheet_exists_only_when_a_chart_is_given_and_holds_the_image(case):
    assert "Graphique" not in workbook(case).sheetnames
    png = X.chart_image(X.raw_table(case["df"]), case["mu"], case["cfg"], "Female", dpi=60)
    wb = workbook(case, chart_png=png)
    assert wb.sheetnames[:5] == ["Table_lissée", "Taux_bruts", "Paramètres", "Critères", "Graphique"]
    assert len(wb["Graphique"]._images) == 1 and wb["Graphique"]._images[0].width == 540


def test_the_multi_method_export_adds_one_sheet_per_tested_method_after_the_others(case):
    png = X.chart_image(X.raw_table(case["df"]), case["mu"], case["cfg"], "Female", dpi=60)
    wb = workbook(case, chart_png=png, methods=X.method_curves(case["fits"]))
    assert wb.sheetnames == ["Table_lissée", "Taux_bruts", "Paramètres", "Critères", "Graphique"] + [f"Méthode_{k}" for k in C.METHODS]
    k = sheet_frame(wb["Méthode_kannisto"])
    assert list(k.columns) == ["age", "mu", "qx"] and k["age"].iloc[0] == 85 and k["age"].iloc[-1] == 110
    np.testing.assert_allclose(k["mu"], case["fits"]["kannisto"].mu_hat, rtol=1e-14)
    few = workbook(case, methods=X.method_curves({"poly3": case["fits"]["poly3"], "beard": case["fits"]["beard"]}))
    assert few.sheetnames[4:] == ["Méthode_poly3", "Méthode_beard"]
    assert all(len(name) <= 31 for name in wb.sheetnames)                                    # Excel's limit on sheet names


def test_final_workbook_assembles_every_sheet_from_the_results_the_app_holds(case):
    data = X.final_workbook(case["df"], case["cfg"], case["fits"], SEL, case["table"], case["mu"], "Female", case["criteria"],
                            columns=["mu", "qx"], include_methods=False, title="Test")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Table_lissée", "Taux_bruts", "Paramètres", "Critères", "Graphique"]
    assert [c.value for c in wb["Table_lissée"][1]] == ["age", "mu", "qx"] and len(wb["Graphique"]._images) == 1
    params = sheet_frame(wb["Paramètres"])
    assert set(params.loc[params["Retenue"], "Clé"]) == set(SEL.values())
    assert len(sheet_frame(wb["Taux_bruts"])) == 101 and len(sheet_frame(wb["Critères"])) == 9
    full = load_workbook(io.BytesIO(X.final_workbook(case["df"], case["cfg"], case["fits"], SEL, case["table"], case["mu"], "Male",
                                                     case["criteria"], include_methods=True)))
    assert full.sheetnames[5:] == [f"Méthode_{k}" for k in C.METHODS] and len(full["Table_lissée"][1]) == 7


# --------------------------------------------------------------------------
# The global spline in the files
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def global_case(case):
    sub = C.zone_slice(case["df"], case["cfg"], "all")
    g = fit_spline_all(sub["age"].to_numpy(), sub["Dx"].to_numpy(), sub["Ex"].to_numpy(), case["cfg"])
    return {**case["fits"], "spline_all": g}


def test_parameter_table_lists_the_global_spline_under_all_ages_and_marks_it_chosen_when_any_zone_took_it(global_case):
    p = X.parameter_table(global_case, SEL)
    assert list(dict.fromkeys(p["Clé"])) == [*C.METHODS, "spline_all"]
    g = p[p["Clé"] == "spline_all"]
    assert set(g["Zone"]) == {"Tous âges"} and set(g["Méthode"]) == {"Spline globale (tous âges)"} and not g["Retenue"].any()
    knots = g[g["Paramètre"].str.startswith("knots[")]
    assert list(knots["Valeur"]) == global_case["spline_all"].params["knots"]
    for zone in C.ZONES:
        pick = {**SEL, zone: "spline_all"}
        q = X.parameter_table(global_case, pick)
        assert q.loc[q["Clé"] == "spline_all", "Retenue"].all()
        assert set(q.loc[q["Retenue"], "Clé"]) == set(pick.values())              # the zone's own pick is no longer chosen


def test_method_curves_end_with_the_global_spline_over_all_the_ages(global_case):
    curves = X.method_curves(global_case)
    assert list(curves) == [*C.METHODS, "spline_all"]
    g = curves["spline_all"]
    assert list(g["age"]) == list(range(101)) and list(g.columns) == ["age", "mu", "qx"]
    np.testing.assert_allclose(g["mu"], global_case["spline_all"].mu_hat, rtol=0)


def test_the_final_workbook_gets_a_sheet_for_the_global_spline_in_the_multi_method_export(case, global_case):
    data = X.final_workbook(case["df"], case["cfg"], global_case, {**SEL, "zone2": "spline_all"}, case["table"], case["mu"],
                            "Female", case["criteria"], include_methods=True)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[5:] == [*(f"Méthode_{k}" for k in C.METHODS), "Méthode_spline_all"]
    params = sheet_frame(wb["Paramètres"])
    assert set(params.loc[params["Retenue"], "Clé"]) == {"poly3", "spline_all", "kannisto"}
